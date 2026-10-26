from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUTDIR = r'C:\Eriel\Codex_Erl\outputs'

def set_font(run, size=10.1, bold=False, color=None, underline=False):
    run.font.name='Times New Roman'; run._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); run._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
    run.font.size=Pt(size); run.bold=bold; run.underline=underline
    if color: run.font.color.rgb=RGBColor(*color)
def add(p,s,size=10.1,bold=False,color=None,underline=False):
    r=p.add_run(s); set_font(r,size,bold,color,underline); return r
def add_rule(p):
    pp=p._p.get_or_add_pPr(); bd=OxmlElement('w:pBdr'); b=OxmlElement('w:bottom')
    b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'10'); b.set(qn('w:space'),'1'); b.set(qn('w:color'),'000000'); bd.append(b); pp.append(bd)
def sec(d,title):
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(5.5); p.paragraph_format.space_after=Pt(1); p.paragraph_format.keep_with_next=True; add(p,title.upper(),12,True); add_rule(p)
def line(d,left,right=None):
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.keep_with_next=True; p.paragraph_format.tab_stops.add_tab_stop(Inches(7.25),WD_TAB_ALIGNMENT.RIGHT); add(p,left,10.75,True)
    if right: p.add_run('\t'); add(p,right,9.7 if len(right) > 30 else 10.75,True)
def para(d,s,size=10.05):
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0; add(p,s,size); return p
def bullet(d,s):
    p=d.add_paragraph(); p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.first_line_indent=Inches(-.12); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0; add(p,'• ',10); add(p,s,10)
def add_skill(d,label,text):
    p=para(d,'',9.95); add(p,label+': ',9.95,True); add(p,text,9.95)

def create(filename, summary, project_title, project_bullets, skills):
    d=Document(); s=d.sections[0]; s.top_margin=Inches(.48); s.bottom_margin=Inches(.42); s.left_margin=Inches(.50); s.right_margin=Inches(.50)
    sty=d.styles['Normal']; sty.font.name='Times New Roman'; sty._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); sty._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman'); sty.font.size=Pt(10.05)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0); add(p,'ERIEL SETIAWAN DEWANTORO',16.5,True)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(5); add(p,'Bekasi, West Java, Indonesia | erieldewantoros@gmail.com | +62 858 8359 1786 | ',9.8); add(p,'LinkedIn',9.8,False,(0,0,255),True); add(p,' | ',9.8); add(p,'GitHub',9.8,False,(0,0,255),True)
    sec(d,'Professional Summary'); para(d,summary)
    sec(d,'Education'); line(d,'Universitas Bina Sarana Informatika','Sep 2022 - Aug 2026'); para(d,'Bachelor of Information Technology',10.6); para(d,'Completed studies and thesis defense in August 2026. Formal graduation scheduled for November 2026.',9.95)
    sec(d,'Selected Data Project'); line(d,project_title,'Undergraduate Thesis');
    for x in project_bullets: bullet(d,x)
    sec(d,'Relevant Experience'); line(d,'Club Tennis Circle Jakarta','Jul 2024 - Nov 2025 | Mar 2026 - Present'); para(d,'Operations & Member Relations',10.6); bullet(d,'Supported member communication, scheduling, registration, operational follow-up, and participant support in a tennis community setting.'); bullet(d,'Coordinated schedule updates and operational information while responding to member inquiries and service-related needs.')
    line(d,'Kementerian Agama Kota Bekasi','Dec 2025 - Feb 2026'); para(d,'Technical Support & Public Relations Intern',10.6); bullet(d,'Supported administrative communication, documentation, internal activity coordination, and basic computer/laptop troubleshooting.')
    sec(d,'Relevant Skills')
    for a,b in skills: add_skill(d,a,b)
    sec(d,'Certifications')
    for x in ['Data Science Bootcamp - Dibimbing.id | 2025 | Final Score: 88.09 | Grade: A','Program Analyst / Analis Program - BNSP | 2025','PCAP: Programming Essentials in Python - Cisco Networking Academy / OpenEDG Python Institute | 2023','Fundamental Database MySQL - Coding Studio | 2024']:
        p=d.add_paragraph(); p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.first_line_indent=Inches(-.12); p.paragraph_format.space_after=Pt(0); add(p,'• ',9.85); add(p,x,9.85,False,(0,0,255),True)
    sec(d,'Languages'); para(d,'Indonesian: Native | English: Intermediate / B1 - able to communicate in daily and professional contexts',9.95)
    d.core_properties.author='Eriel Setiawan Dewantoro'; d.core_properties.title=filename.replace('_',' '); d.save(OUTDIR+'\\'+filename)

thesis = [
    'Developed a binary classification approach for Jakarta ISPU conditions, classifying Tidak Sehat+ versus Bukan Tidak Sehat+ using 2020-2024 data.',
    'Prepared and analyzed 1,826 valid observations; compared Random Forest, XGBoost, LSTM, and GRU using a chronological train-test split.',
    'Selected Random Forest based on the strongest overall test performance: F1-score 0.9600 and ROC-AUC 0.9813.'
]
create('CV_Eriel_Setiawan_Dewantoro_Business_Intelligence_tfam.docx',
       'Information Technology graduate candidate with a foundation in SQL, Python, Pandas, data cleaning, exploratory data analysis, reporting, and data visualization. Completed an undergraduate data science thesis using a chronological evaluation approach and machine learning models, and holds a Data Science Bootcamp certification. Comfortable working with CSV/XLSX data, Microsoft Excel, Google Sheets, MySQL, and DBeaver. Interested in contributing to entry-level Business Intelligence reporting and analysis in a hybrid environment.',
       'ISPU Jakarta Classification', thesis,
       [('SQL & Data', 'SQL, MySQL, DBeaver, CSV/XLSX, data cleaning, data validation, exploratory data analysis.'),('Analysis & Reporting', 'Python, Pandas, Google Colab, Microsoft Excel, Google Sheets, reporting, data visualization, dashboard fundamentals.'),('Technical', 'Java, GitHub, basic database concepts, HTML, CSS, JavaScript.'),('Professional', 'Analytical thinking, attention to detail, problem solving, documentation, stakeholder communication.')])
create('CV_Eriel_Setiawan_Dewantoro_Data_Analyst.docx',
       'Information Technology graduate candidate with hands-on academic and project experience in SQL, Python, Pandas, data cleaning, exploratory data analysis, reporting, and data visualization. Completed an undergraduate thesis that prepared and analyzed Jakarta ISPU data and compared machine learning models using a chronological test set. Detail-oriented and comfortable translating datasets into structured findings through clear documentation, reporting, and communication. Interested in an entry-level Data Analyst role where analytical work supports practical business decisions.',
       'ISPU Jakarta Classification', thesis,
       [('Data Analysis', 'SQL, Python, Pandas, Google Colab, data cleaning, data validation, exploratory data analysis, descriptive analysis.'),('Reporting & Visualization', 'Microsoft Excel, Google Sheets, CSV/XLSX, reporting, data visualization, dashboard fundamentals.'),('Database & Technical', 'MySQL, DBeaver, Java, GitHub, basic database concepts.'),('Professional', 'Analytical thinking, problem solving, attention to detail, documentation, written and verbal communication.')])
print('Created two DOCX files')
