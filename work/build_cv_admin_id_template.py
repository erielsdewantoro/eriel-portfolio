from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r'C:\Eriel\Codex_Erl\outputs\CV_Eriel_Staff_Admin_Data_Entry_ID.docx'

def font(run, size=10.25, bold=False, color=None, underline=False):
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    run.font.size = Pt(size); run.bold = bold; run.underline = underline
    if color: run.font.color.rgb = RGBColor(*color)
def txt(p, s, size=10.25, bold=False, color=None, underline=False):
    r=p.add_run(s); font(r,size,bold,color,underline); return r
def rule(p):
    pPr=p._p.get_or_add_pPr(); b=OxmlElement('w:pBdr'); x=OxmlElement('w:bottom')
    x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'10'); x.set(qn('w:space'),'1'); x.set(qn('w:color'),'000000'); b.append(x); pPr.append(b)
def section(doc, title):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(1); p.paragraph_format.keep_with_next=True
    txt(p,title.upper(),12,True); rule(p)
def line(doc,left,right=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.keep_with_next=True; p.paragraph_format.tab_stops.add_tab_stop(Inches(7.25),WD_TAB_ALIGNMENT.RIGHT)
    txt(p,left,10.8,True)
    if right:
        p.add_run('\t')
        txt(p,right,9.7 if len(right) > 30 else 10.8,True)
def para(doc,s,size=10.15):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0; txt(p,s,size); return p
def bullet(doc,s):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.first_line_indent=Inches(-.12); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
    txt(p,'• ',10.1); txt(p,s,10.1)

doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.48); sec.bottom_margin=Inches(.42); sec.left_margin=Inches(.50); sec.right_margin=Inches(.50)
base=doc.styles['Normal']; base.font.name='Times New Roman'; base._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); base._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman'); base.font.size=Pt(10.15)

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0); txt(p,'ERIEL SETIAWAN DEWANTORO',16.5,True)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(5)
txt(p,'Bekasi, Jawa Barat, Indonesia | erieldewantoros@gmail.com | +62 858 8359 1786 | ',9.9); txt(p,'LinkedIn',9.9,False,(0,0,255),True); txt(p,' | ',9.9); txt(p,'GitHub',9.9,False,(0,0,255),True)

section(doc,'Profil Profesional')
para(doc,'Kandidat fresh graduate Teknologi Informasi yang telah menyelesaikan studi dan sidang skripsi pada Agustus 2026, dengan pengalaman dalam dukungan administrasi, koordinasi operasional, komunikasi anggota, registrasi, dan dokumentasi. Terbiasa menggunakan Microsoft Excel, Word, PowerPoint, Google Sheets, Google Workspace, dan Google Earth untuk data entry, mapping area, dan kebutuhan administrasi harian. Teliti, responsif, memiliki laptop pribadi, dan siap bekerja di area Cinere.')

section(doc,'Pendidikan')
line(doc,'Universitas Bina Sarana Informatika','Sep 2022 - Agu 2026')
para(doc,'S1 Teknologi Informasi',10.65)
para(doc,'Studi dan sidang skripsi selesai pada Agustus 2026. Wisuda/formal graduation dijadwalkan pada November 2026.',10.0)

section(doc,'Pengalaman Kerja')
line(doc,'Club Tennis Circle Jakarta','Jul 2024 - Nov 2025 | Mar 2026 - Present')
para(doc,'Operations & Member Relations',10.65)
bullet(doc,'Mendukung operasional komunitas sehari-hari, termasuk komunikasi anggota, penjadwalan, registrasi, dan dukungan peserta.')
bullet(doc,'Mengoordinasikan pembaruan jadwal dan informasi operasional untuk membantu menjaga komunikasi yang akurat serta pengalaman anggota yang positif.')
bullet(doc,'Menangani pertanyaan dan kebutuhan layanan anggota secara responsif, termasuk melakukan follow-up atas kebutuhan operasional.')
bullet(doc,'Mendukung aktivitas tenis dan matchday antar-komunitas melalui koordinasi peserta serta bantuan operasional on-site.')
line(doc,'Kementerian Agama Kota Bekasi','Des 2025 - Feb 2026')
para(doc,'Technical Support & Public Relations Intern',10.65)
bullet(doc,'Mendukung komunikasi administratif, dokumentasi, serta persiapan kegiatan internal dan kebutuhan organisasi.')
bullet(doc,'Membantu koordinasi dan kegiatan kepanitiaan internal untuk mendukung kelancaran acara serta operasional sehari-hari.')
bullet(doc,'Membuat materi visual dan komunikasi menggunakan Canva, Microsoft Word, dan CapCut, serta memberikan troubleshooting dasar komputer dan laptop.')

section(doc,'Keahlian Relevan')
p=para(doc,'',10.0); txt(p,'Administrasi & Data Entry: ',10.0,True); txt(p,'Microsoft Excel, Microsoft Word, Microsoft PowerPoint, Google Sheets, Google Workspace, data entry, dokumentasi, administrative support.',10.0)
p=para(doc,'',10.0); txt(p,'Mapping & Operasional: ',10.0,True); txt(p,'Google Earth, mapping area, penjadwalan, registrasi, koordinasi operasional, participant support, follow-up.',10.0)
p=para(doc,'',10.0); txt(p,'Teknis & Kesiapan Kerja: ',10.0,True); txt(p,'Troubleshooting dasar komputer/laptop, laptop pribadi, siap bekerja di Cinere.',10.0)
p=para(doc,'',10.0); txt(p,'Soft Skills: ',10.0,True); txt(p,'Komunikasi, ketelitian, problem solving, adaptabilitas, stakeholder communication, kerja sama tim, time management.',10.0)

section(doc,'Sertifikasi')
for s in ['Data Science Bootcamp - Dibimbing.id | 2025 | Final Score: 88.09 | Grade: A','Program Analyst / Analis Program - BNSP | 2025','PCAP: Programming Essentials in Python - Cisco Networking Academy / OpenEDG Python Institute | 2023','Fundamental Database MySQL - Coding Studio | 2024']:
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.first_line_indent=Inches(-.12); p.paragraph_format.space_after=Pt(0); txt(p,'• ',10.0); txt(p,s,10.0,False,(0,0,255),True)
section(doc,'Bahasa')
para(doc,'Bahasa Indonesia: Native | Bahasa Inggris: Intermediate / B1 - dapat digunakan untuk komunikasi sehari-hari dan konteks profesional',10.0)
doc.core_properties.author='Eriel Setiawan Dewantoro'; doc.core_properties.title='CV Eriel Staff Admin Data Entry Bahasa Indonesia'; doc.save(OUT); print(OUT)
