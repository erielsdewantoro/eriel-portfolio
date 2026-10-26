from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

OUT = r"C:\Eriel\Codex_Erl\outputs\CV_Eriel_Staff_Admin_Data_Entry.docx"

def set_font(run, name='Aptos', size=10, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for side, val in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            tcMar.append(node)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')

def add_text(p, text, size=10, bold=False, color=None):
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color)
    return r

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.first_line_indent = Inches(-0.13)
    add_text(p, text, 9.4)

def section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    add_text(p, text.upper(), 10.1, True)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '2')
    bottom.set(qn('w:color'), '1F4E79')
    pbdr.append(bottom)
    pPr.append(pbdr)

def experience(doc, title, org, dates, location, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(1)
    add_text(p, title, 10, True)
    p.add_run('  |  ')
    add_text(p, org, 10, False)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(1.5)
    add_text(p2, f'{dates} | {location}', 8.8, False, (80,80,80))
    for item in bullets:
        add_bullet(doc, item)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.48)
sec.bottom_margin = Inches(0.45)
sec.left_margin = Inches(0.6)
sec.right_margin = Inches(0.6)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].font.size = Pt(10)
styles['Normal'].paragraph_format.space_after = Pt(0)

# Header
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_text(p, 'ERIEL SETIAWAN DEWANTORO', 18, True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
add_text(p, 'Staff Admin | Data Entry | Operational Support', 10.5, True, (31, 78, 121))
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(5)
add_text(p, 'Bekasi, Jawa Barat | +62 858 8359 1786 | erieldewantoros@gmail.com', 9.2)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
add_text(p, 'linkedin.com/in/eriel-setiawan-dewantoro | github.com/erielsdewantoro', 8.7, False, (65,65,65))

section_heading(doc, 'Profil')
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
add_text(p, 'Kandidat dengan pendidikan S1 Teknologi Informasi selesai Agustus 2026, berpengalaman mendukung administrasi, komunikasi operasional, koordinasi jadwal, registrasi, dokumentasi, dan kebutuhan anggota komunitas. Terbiasa bekerja secara rapi dan responsif dalam aktivitas yang melibatkan data, follow-up, serta koordinasi dengan banyak pihak. Memiliki kemampuan Microsoft Excel, Microsoft Word, Microsoft PowerPoint, Google Sheets, Google Earth, dan Google Workspace untuk mendukung kebutuhan administrasi, data entry, dan mapping area.', 9.5)

section_heading(doc, 'Kompetensi Relevan')
table = doc.add_table(rows=2, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
skills = [
    ('Administrasi & Data Entry', 'Microsoft Excel, Microsoft Word, Microsoft PowerPoint, Google Sheets, Google Workspace'),
    ('Koordinasi Operasional', 'Penjadwalan, registrasi, follow-up, participant support'),
    ('Dokumentasi & Komunikasi', 'Dokumentasi kegiatan, komunikasi administratif, stakeholder communication'),
    ('Mapping & Presentasi', 'Google Earth, Microsoft PowerPoint'),
    ('Teknis', 'Basic computer/laptop troubleshooting, CSV/XLSX, GitHub'),
    ('Kesiapan Kerja', 'Memiliki laptop pribadi dan siap bekerja di area Cinere')
]
for i, (head, body) in enumerate(skills):
    cell = table.cell(i // 3, i % 3)
    cell.width = Inches(2.4)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_margins(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    add_text(p, head, 8.7, True)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    add_text(p, body, 8.5)
    if i % 2 == 0:
        set_cell_shading(cell, 'F3F6F9')

section_heading(doc, 'Pengalaman Kerja')
experience(doc, 'Operations & Member Relations', 'Club Tennis Circle Jakarta', 'Juli 2024 - Agustus 2026', 'Jakarta, Indonesia', [
    'Mendukung operasional komunitas sehari-hari, termasuk komunikasi anggota, penjadwalan, registrasi, dan dukungan peserta.',
    'Mengomunikasikan pembaruan jadwal serta informasi operasional untuk membantu kelancaran koordinasi dan pengalaman anggota.',
    'Menangani pertanyaan dan kebutuhan layanan anggota secara responsif serta melakukan follow-up sesuai kebutuhan.',
    'Mendukung kegiatan tenis dan matchday antar-komunitas melalui koordinasi peserta, dukungan on-site, dan tindak lanjut operasional.'
])
experience(doc, 'Technical Support & Public Relations Intern', 'Kementerian Agama Kota Bekasi', 'Desember 2025 - Februari 2026', 'Bekasi, Indonesia', [
    'Mendukung komunikasi public relations dan administrasi untuk kegiatan internal serta kebutuhan organisasi.',
    'Membantu persiapan kegiatan, koordinasi panitia, dan dokumentasi untuk mendukung kelancaran aktivitas internal.',
    'Membuat dan mengedit materi visual serta komunikasi menggunakan Canva, Microsoft Word, dan CapCut.',
    'Memberikan dukungan troubleshooting dasar komputer dan laptop untuk kebutuhan operasional sehari-hari.'
])

section_heading(doc, 'Pendidikan')
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
add_text(p, 'S1 Teknologi Informasi', 10, True)
add_text(p, '  |  Universitas Bina Sarana Informatika', 10)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
add_text(p, 'September 2022 - Agustus 2026 | Studi dan sidang skripsi selesai Agustus 2026 | Wisuda direncanakan November 2026', 8.8, False, (80,80,80))

section_heading(doc, 'Sertifikasi')
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
add_text(p, 'Data Science Bootcamp - Dibimbing.id (2025, Final Score 88.09, Grade A)  |  Program Analyst / Analis Program - BNSP (2025)  |  Fundamental Database MySQL - Coding Studio (2024)  |  PCAP Programming Essentials in Python - Cisco Networking Academy / OpenEDG Python Institute (2023)', 8.75)

section_heading(doc, 'Bahasa')
p = doc.add_paragraph()
add_text(p, 'Bahasa Indonesia: Native  |  Bahasa Inggris: Intermediate (B1) - komunikasi sehari-hari dan konteks profesional', 9)

# Compact footer
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_p.paragraph_format.space_before = Pt(3)
add_text(footer_p, 'CV tailored for Staff Admin - Data Entry and Area Mapping', 7.5, False, (110,110,110))

doc.core_properties.author = 'Eriel Setiawan Dewantoro'
doc.core_properties.title = 'CV Eriel Setiawan Dewantoro Staff Admin Data Entry'
doc.save(OUT)
print(OUT)
