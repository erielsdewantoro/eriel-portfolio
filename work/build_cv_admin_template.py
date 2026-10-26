from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r'C:\Eriel\Codex_Erl\outputs\CV_Eriel_Staff_Admin_Data_Entry_Revision.docx'

def set_font(run, size=10.4, bold=False, color=None, underline=False):
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    run.font.size = Pt(size)
    run.bold = bold
    run.underline = underline
    if color:
        run.font.color.rgb = RGBColor(*color)

def add_text(p, text, size=10.4, bold=False, color=None, underline=False):
    r = p.add_run(text)
    set_font(r, size, bold, color, underline)
    return r

def set_bottom_rule(p):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '10')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '000000')
    pBdr.append(bottom)
    pPr.append(pBdr)

def section(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    add_text(p, title.upper(), 12, True)
    set_bottom_rule(p)

def item_line(doc, left, right=None, bold=True, size=10.8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.25), WD_TAB_ALIGNMENT.RIGHT)
    add_text(p, left, size, bold)
    if right:
        p.add_run('\t')
        add_text(p, right, size, True)
    return p

def body(doc, text, size=10.25, after=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    add_text(p, text, size)
    return p

def bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    add_text(p, '• ', 10.25)
    add_text(p, text, 10.25)
    return p

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.48)
sec.bottom_margin = Inches(0.42)
sec.left_margin = Inches(0.50)
sec.right_margin = Inches(0.50)

normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
normal.font.size = Pt(10.25)
normal.paragraph_format.space_after = Pt(0)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(0)
add_text(p, 'ERIEL SETIAWAN DEWANTORO', 16.5, True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(5)
add_text(p, 'Bekasi, West Java, Indonesia | erieldewantoros@gmail.com | +62 858 8359 1786 | ', 9.9)
add_text(p, 'LinkedIn', 9.9, False, (0,0,255), True)
add_text(p, ' | ', 9.9)
add_text(p, 'GitHub', 9.9, False, (0,0,255), True)

section(doc, 'Professional Summary')
body(doc, 'Information Technology graduate candidate with completed studies and thesis defense in August 2026, with hands-on experience in administrative support, operational coordination, member communication, registration, and documentation. Comfortable using Microsoft Excel, Word, PowerPoint, Google Sheets, Google Workspace, and Google Earth for data entry, area mapping, and day-to-day administrative needs. Detail-oriented, responsive, and ready to work in Cinere with a personal laptop.', 10.25)

section(doc, 'Education')
item_line(doc, 'Universitas Bina Sarana Informatika', 'Sep 2022 - Aug 2026')
body(doc, 'Bachelor of Information Technology', 10.7)
body(doc, 'Completed studies and thesis defense in August 2026. Formal graduation scheduled for November 2026.', 10.05)

section(doc, 'Work Experience')
item_line(doc, 'Club Tennis Circle Jakarta', 'Jul 2024 - Aug 2026')
body(doc, 'Operations & Member Relations', 10.7, 0)
bullet(doc, 'Supported day-to-day community operations, including member communication, scheduling, registration, and participant support.')
bullet(doc, 'Coordinated schedule updates and operational information to help maintain accurate communication and a positive member experience.')
bullet(doc, 'Handled member inquiries and service-related needs responsively, including follow-up on operational requirements.')
bullet(doc, 'Assisted tennis activities and inter-community matchday support through participant coordination and on-site operational assistance.')

item_line(doc, 'Kementerian Agama Kota Bekasi', 'Dec 2025 - Feb 2026')
body(doc, 'Technical Support & Public Relations Intern', 10.7, 0)
bullet(doc, 'Supported administrative communication, documentation, and event preparation for internal activities and organizational needs.')
bullet(doc, 'Assisted coordination and internal committee activities to support the smooth execution of events and daily operations.')
bullet(doc, 'Created visual and communication materials using Canva, Microsoft Word, and CapCut, and provided basic computer and laptop troubleshooting support.')

section(doc, 'Relevant Skills')
p = body(doc, '', 10.05)
add_text(p, 'Administrative & Data Entry: ', 10.05, True)
add_text(p, 'Microsoft Excel, Microsoft Word, Microsoft PowerPoint, Google Sheets, Google Workspace, data entry, documentation, administrative support.', 10.05)
p = body(doc, '', 10.05)
add_text(p, 'Mapping & Operations: ', 10.05, True)
add_text(p, 'Google Earth, area mapping, scheduling, registration, operational coordination, participant support, follow-up.', 10.05)
p = body(doc, '', 10.05)
add_text(p, 'Technical & Work Readiness: ', 10.05, True)
add_text(p, 'Basic computer and laptop troubleshooting, personal laptop, ready to work in Cinere.', 10.05)
p = body(doc, '', 10.05)
add_text(p, 'Soft Skills: ', 10.05, True)
add_text(p, 'Communication, attention to detail, problem solving, adaptability, stakeholder communication, teamwork, time management.', 10.05)

section(doc, 'Certifications')
for text in [
    'Data Science Bootcamp - Dibimbing.id | 2025 | Final Score: 88.09 | Grade: A',
    'Program Analyst / Analis Program - BNSP | 2025',
    'PCAP: Programming Essentials in Python - Cisco Networking Academy / OpenEDG Python Institute | 2023',
    'Fundamental Database MySQL - Coding Studio | 2024',
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(0)
    add_text(p, '• ', 10.0)
    add_text(p, text, 10.0, False, (0,0,255), True)

section(doc, 'Languages')
body(doc, 'Indonesian: Native | English: Intermediate / B1 - able to communicate in daily and professional contexts', 10.05)

doc.core_properties.author = 'Eriel Setiawan Dewantoro'
doc.core_properties.title = 'CV Eriel Setiawan Dewantoro Staff Admin Data Entry'
doc.save(OUT)
print(OUT)
